"""
tests/test_exports.py
=====================
Unit tests for core/exports.py (Excel workbook and Markdown generation).
"""

from __future__ import annotations

import io
import pandas as pd
import pytest
from openpyxl import load_workbook

from core.exports import export_excel, export_markdown_summary
from core.demo_data import generate_demo_findings


@pytest.fixture
def sample_df():
    findings = generate_demo_findings(target_count=30, seed=42)
    return pd.DataFrame(findings)


def test_export_excel_creates_valid_workbook(sample_df):
    excel_bytes = export_excel(sample_df)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Summary" in sheet_names
    assert "Register" in sheet_names
    assert "Overdue" in sheet_names
    assert "Charts" in sheet_names

    # Check rows exist in Register
    ws_reg = wb["Register"]
    assert ws_reg.max_row > 1


def test_export_markdown_summary(sample_df):
    md = export_markdown_summary(sample_df)
    assert isinstance(md, str)
    assert "# Weekly Vulnerability Management Summary" in md
    assert "Executive Overview" in md
    assert "Top 5 Highest Risk Open Findings" in md
    assert "Remediation Progress" in md
    assert "Recommended Actions" in md
